from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

from docx import Document as DocxDocument
from reportlab.lib.pagesizes import A4
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.config import get_settings
from app.db.models import ExportJob, Flashcard, FlashcardDeck, Question, QuestionOption, Quiz, StudyPlan, StudyTask

settings = get_settings()


def _resource_lines(db: Session, resource_type: str, resource_id: int, user_id: int) -> tuple[str, list[str]]:
    if resource_type == "QUIZ":
        quiz = db.get(Quiz, resource_id)
        if not quiz or quiz.owner_id != user_id:
            raise ValueError("Quiz not found")
        lines = []
        questions = db.scalars(select(Question).where(Question.quiz_id == quiz.id).order_by(Question.question_order)).all()
        for q in questions:
            lines.append(f"Câu {q.question_order}: {q.question_text}")
            for o in db.scalars(select(QuestionOption).where(QuestionOption.question_id == q.id).order_by(QuestionOption.position)).all():
                marker = " *" if o.is_correct else ""
                lines.append(f"  {o.option_key}. {o.option_text}{marker}")
            if q.explanation:
                lines.append(f"  Giải thích: {q.explanation}")
            lines.append("")
        return quiz.title, lines
    if resource_type == "FLASHCARD_DECK":
        deck = db.get(FlashcardDeck, resource_id)
        if not deck or deck.owner_id != user_id:
            raise ValueError("Flashcard deck not found")
        cards = db.scalars(select(Flashcard).where(Flashcard.deck_id == deck.id).order_by(Flashcard.card_order)).all()
        lines = [f"{c.card_order}. {c.front_text}\n   → {c.back_text}" for c in cards]
        return deck.title, lines
    if resource_type == "STUDY_PLAN":
        plan = db.get(StudyPlan, resource_id)
        if not plan or plan.user_id != user_id:
            raise ValueError("Study plan not found")
        tasks = db.scalars(select(StudyTask).where(StudyTask.plan_id == plan.id).order_by(StudyTask.task_date, StudyTask.sort_order)).all()
        lines = [f"{t.task_date}: [{t.task_type}] {t.title} ({t.estimated_minutes} phút)" for t in tasks]
        return plan.title, lines
    raise ValueError("Unsupported resource type")


def create_export(db: Session, user_id: int, resource_type: str, resource_id: int, file_format: str) -> ExportJob:
    job = ExportJob(user_id=user_id, resource_type=resource_type, resource_id=resource_id, file_format=file_format, status="PROCESSING")
    db.add(job); db.commit(); db.refresh(job)
    try:
        title, lines = _resource_lines(db, resource_type, resource_id, user_id)
        out_dir = Path("storage/exports") / str(user_id)
        out_dir.mkdir(parents=True, exist_ok=True)
        suffix = ".docx" if file_format == "DOCX" else ".pdf"
        path = out_dir / f"{resource_type.lower()}_{resource_id}_{job.id}{suffix}"
        if file_format == "DOCX":
            doc = DocxDocument()
            doc.add_heading(title, level=1)
            for line in lines:
                doc.add_paragraph(line)
            doc.save(path)
        else:
            styles = getSampleStyleSheet()
            story = [Paragraph(title, styles["Title"]), Spacer(1, 12)]
            for line in lines:
                story.append(Paragraph(line.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"), styles["BodyText"]))
                story.append(Spacer(1, 6))
            SimpleDocTemplate(str(path), pagesize=A4).build(story)
        job.status = "COMPLETED"
        job.file_url = str(path.resolve())
        job.completed_at = datetime.now(timezone.utc)
        db.commit(); db.refresh(job)
        return job
    except Exception as exc:
        job.status = "FAILED"
        job.error_message = str(exc)
        job.completed_at = datetime.now(timezone.utc)
        db.commit()
        raise
